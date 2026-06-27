"""One-time script to build templates/custom_starter/starter_template.docx."""

from __future__ import annotations

from pathlib import Path

from docx import Document

# Jinja placeholders — must match resume_json keys used by docx_renderer.py / custom_docx_renderer.py
STARTER_LINES = [
    "{{ name }}",
    "{{ contact.email }} | {{ contact.phone }}",
    "",
    "SUMMARY",
    "{{ summary }}",
    "",
    "EXPERIENCE",
    "{% for exp in experience %}",
    "{{ exp.title }} — {{ exp.company }} ({{ exp.dates }})",
    "{% for bullet in exp.bullets %}",
    "• {{ bullet }}",
    "{% endfor %}",
    "{% endfor %}",
    "",
    "EDUCATION",
    "{% for edu in education %}",
    "{{ edu.degree }}, {{ edu.institution }} ({{ edu.dates }})",
    "{% endfor %}",
    "",
    "SKILLS",
    '{{ skills | join(", ") }}',
]


def main() -> None:
    out_dir = Path(__file__).resolve().parent.parent / "templates" / "custom_starter"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "starter_template.docx"

    doc = Document()
    for line in STARTER_LINES:
        doc.add_paragraph(line)

    doc.save(out_path)
    print(f"Wrote {out_path} ({out_path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
